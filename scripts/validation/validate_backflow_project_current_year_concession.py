from __future__ import annotations

import csv
import json
import sys
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator, Sequence

from openpyxl import load_workbook

from _project_root import find_project_root


ROOT = find_project_root(__file__)
REPORT_MONTH = sys.argv[1] if len(sys.argv) > 1 else "202608"
REPORT_YEAR = REPORT_MONTH[:4]
TOLERANCE_WAN = 1e-6  # 0.01 yuan
YUAN_PER_WAN = 10_000.0


def u(text: str) -> str:
    return text.encode("utf-8").decode("unicode_escape")


REPORT_TOKEN = u(r"\u56de\u6b3e\u8425\u6536\u6bd4")
PROJECT_TOKEN = u(r"\u9879\u76ee")
LEDGER_TOKEN = u(r"\u7efc\u7ba1\u6298\u8ba9\u8c03\u6574\u9879\u53f0\u8d26")
QUERY_TOKEN = u(r"\u9879\u76ee\u67e5\u8be2")

CODE = u(r"\u7acb\u9879\u7f16\u7801")
PROJECT_NAME = u(r"\u9879") + u(r"\u76ee")
REGION = u(r"\u533a\u57df")
LINE = u(r"\u6761\u7ebf")
METRIC = u(r"\u5f53\u5e74\u5e94\u6536\u6b3e\u7efc\u7ba1\u6298\u8ba9\u91d1\u989d")
DATA_MONTH = u(r"\u6570\u636e\u5e74\u6708")
LEVEL_ONE_TYPE = u(r"\u4e00\u7ea7\u6298\u8ba9\u7c7b\u578b")
CURRENT_YEAR_AMOUNT_YUAN = u(
    r"\u5176\u4e2d\uff1a\u5f53\u5e74\u5e94\u6536\u6b3e\u6298\u8ba9\u91d1\u989d\uff08\u5355\u4f4d\uff1a\u5143\uff09"
)
ELIGIBLE_TYPES = {
    u(r"\u56de\u6b3e\u8425\u6536\u6bd4"),
    u(r"\u901a\u7528"),
}


@dataclass
class LedgerMonth:
    month: str
    path: Path
    amount_by_code_yuan: dict[str, float]
    record_count_by_code: Counter[str]
    eligible_rows: int
    exact_duplicate_rows: int
    blank_code_rows: int
    invalid_amount_rows: int


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def normalize_code(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    if text.endswith(".0"):
        text = text[:-2]
    if text in {"", "NAN", "NONE"}:
        return ""
    return text


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_month(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y%m")
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    return digits[:6]


def numeric_or_none(value: object) -> float | None:
    if value is None or clean_text(value) == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def near_zero(value: float) -> float:
    return 0.0 if abs(value) < TOLERANCE_WAN else value


def workbook_rows(path: Path) -> Iterator[tuple[object, ...]]:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        # These system exports incorrectly declare the used range as A1.  Resetting
        # the dimensions makes openpyxl stream the actual worksheet XML rows.
        sheet.reset_dimensions()
        yield from sheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def find_unique_workbook(token: str, month: str | None = None, suffix_token: str | None = None) -> Path:
    matches = []
    for path in ROOT.iterdir():
        if not path.is_file() or path.suffix.lower() != ".xlsx":
            continue
        if token not in path.stem:
            continue
        if month is not None and month not in path.stem:
            continue
        if suffix_token is not None and not path.stem.endswith(suffix_token):
            continue
        matches.append(path)
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one workbook for token={token!r}, month={month!r}, suffix={suffix_token!r}; "
            f"got {len(matches)}: {[path.name for path in matches]}"
        )
    return matches[0]


def header_index(row: Sequence[object]) -> dict[str, int]:
    return {clean_text(value): idx for idx, value in enumerate(row) if clean_text(value)}


def require_headers(index: dict[str, int], required: Iterable[str], path: Path) -> None:
    missing = [label for label in required if label not in index]
    if missing:
        raise RuntimeError(f"Missing headers in {path.name}: {missing}")


def value_at(row: Sequence[object], index: int) -> object:
    return row[index] if index < len(row) else None


def load_report(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    rows = workbook_rows(path)
    header: dict[str, int] | None = None
    key_header: dict[str, int] | None = None
    report_rows: list[dict[str, object]] = []
    for row_number, row in enumerate(rows, start=1):
        candidate = header_index(row)
        header_completed = False
        if header is None and CODE in candidate:
            key_header = candidate
            if METRIC in candidate:
                header = candidate
                header_completed = True
        elif header is None and key_header is not None and METRIC in candidate:
            # The report uses a two-tier header: identifiers are on the first
            # header row while metric names are on the next row.
            header = {**key_header, **candidate}
            header_completed = True
        if header_completed:
            require_headers(header, [CODE, PROJECT_NAME, REGION, LINE, METRIC], path)
            continue
        if header is None:
            continue
        code = normalize_code(value_at(row, header[CODE]))
        if not code:
            continue
        raw_metric = value_at(row, header[METRIC])
        report_rows.append(
            {
                "report_row": row_number,
                "project_code": code,
                "project_name": clean_text(value_at(row, header[PROJECT_NAME])),
                "region": clean_text(value_at(row, header[REGION])),
                "line": clean_text(value_at(row, header[LINE])),
                "report_raw": raw_metric,
                "report_is_blank": raw_metric is None or clean_text(raw_metric) == "",
                "report_value_wan": numeric_or_none(raw_metric),
            }
        )
    if header is None:
        raise RuntimeError(f"Could not find report header in {path.name}")

    code_counts = Counter(str(row["project_code"]) for row in report_rows)
    diagnostics = {
        "rows": len(report_rows),
        "unique_codes": len(code_counts),
        "duplicate_code_count": sum(1 for count in code_counts.values() if count > 1),
        "duplicate_row_count": sum(count - 1 for count in code_counts.values() if count > 1),
        "blank_metric_rows": sum(bool(row["report_is_blank"]) for row in report_rows),
        "non_numeric_metric_rows": sum(
            row["report_value_wan"] is None and not bool(row["report_is_blank"])
            for row in report_rows
        ),
    }
    return report_rows, diagnostics


def load_ledger(path: Path, expected_month: str) -> LedgerMonth:
    header: dict[str, int] | None = None
    sums: defaultdict[str, float] = defaultdict(float)
    counts: Counter[str] = Counter()
    duplicate_signatures: Counter[tuple[object, ...]] = Counter()
    eligible_rows = 0
    blank_code_rows = 0
    invalid_amount_rows = 0

    required = [CODE, DATA_MONTH, LEVEL_ONE_TYPE, CURRENT_YEAR_AMOUNT_YUAN]
    for row in workbook_rows(path):
        candidate = header_index(row)
        if header is None and CODE in candidate and LEVEL_ONE_TYPE in candidate:
            header = candidate
            require_headers(header, required, path)
            continue
        if header is None:
            continue

        period = normalize_month(value_at(row, header[DATA_MONTH]))
        ledger_type = clean_text(value_at(row, header[LEVEL_ONE_TYPE]))
        if period != expected_month or ledger_type not in ELIGIBLE_TYPES:
            continue

        eligible_rows += 1
        code = normalize_code(value_at(row, header[CODE]))
        if not code:
            blank_code_rows += 1
            continue

        raw_amount = value_at(row, header[CURRENT_YEAR_AMOUNT_YUAN])
        amount = numeric_or_none(raw_amount)
        if amount is None:
            invalid_amount_rows += 1
            amount = 0.0

        sums[code] += amount
        counts[code] += 1
        signature = tuple(value_at(row, idx) for idx in range(len(row)))
        duplicate_signatures[signature] += 1

    if header is None:
        raise RuntimeError(f"Could not find ledger header in {path.name}")

    return LedgerMonth(
        month=expected_month,
        path=path,
        amount_by_code_yuan=dict(sums),
        record_count_by_code=counts,
        eligible_rows=eligible_rows,
        exact_duplicate_rows=sum(count - 1 for count in duplicate_signatures.values() if count > 1),
        blank_code_rows=blank_code_rows,
        invalid_amount_rows=invalid_amount_rows,
    )


def load_query_codes(path: Path) -> tuple[set[str], dict[str, object]]:
    header: dict[str, int] | None = None
    codes: Counter[str] = Counter()
    blank_code_rows = 0
    for row in workbook_rows(path):
        candidate = header_index(row)
        if header is None and CODE in candidate:
            header = candidate
            continue
        if header is None:
            continue
        code = normalize_code(value_at(row, header[CODE]))
        if code:
            codes[code] += 1
        elif any(clean_text(value) for value in row):
            blank_code_rows += 1
    if header is None:
        raise RuntimeError(f"Could not find query header in {path.name}")
    return set(codes), {
        "unique_codes": len(codes),
        "duplicate_code_count": sum(1 for count in codes.values() if count > 1),
        "duplicate_row_count": sum(count - 1 for count in codes.values() if count > 1),
        "blank_code_rows": blank_code_rows,
    }


def month_range(report_month: str) -> list[str]:
    month = int(report_month[4:])
    if not 1 <= month <= 12:
        raise RuntimeError(f"Invalid report month: {report_month!r}")
    return [f"{report_month[:4]}{item:02d}" for item in range(1, month + 1)]


def candidate_summary(details: list[dict[str, object]], prefix: str) -> dict[str, object]:
    numeric_rows = [row for row in details if row["report_value_wan"] is not None]
    diff_key = f"diff_{prefix}_wan"
    match_key = f"match_{prefix}_numeric"
    presence_key = f"match_{prefix}_presence"
    diffs = [float(row[diff_key]) for row in numeric_rows]
    mismatches = [row for row in numeric_rows if not bool(row[match_key])]
    presence_mismatches = [row for row in details if not bool(row[presence_key])]
    report_total = sum(float(row["report_value_wan"] or 0.0) for row in details)
    calc_total = sum(float(row[f"calc_{prefix}_wan"]) for row in details)
    return {
        "numeric_rows": len(numeric_rows),
        "numeric_match_rows": sum(bool(row[match_key]) for row in numeric_rows),
        "numeric_mismatch_rows": len(mismatches),
        "presence_match_rows": sum(bool(row[presence_key]) for row in details),
        "presence_mismatch_rows": len(presence_mismatches),
        "max_abs_diff_wan": max((abs(value) for value in diffs), default=0.0),
        "report_total_wan": report_total,
        "calculated_total_wan": calc_total,
        "difference_total_wan": report_total - calc_total,
    }


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    configure_stdout()
    report_path = find_unique_workbook(REPORT_TOKEN, REPORT_MONTH, PROJECT_TOKEN)
    query_path = find_unique_workbook(QUERY_TOKEN)
    report_rows, report_diagnostics = load_report(report_path)
    query_codes, query_diagnostics = load_query_codes(query_path)

    ledgers: list[LedgerMonth] = []
    for month in month_range(REPORT_MONTH):
        ledger_path = find_unique_workbook(LEDGER_TOKEN, month)
        ledgers.append(load_ledger(ledger_path, month))

    report_codes = {str(row["project_code"]) for row in report_rows}
    current_ledger = ledgers[-1]
    ytd_amount_yuan: defaultdict[str, float] = defaultdict(float)
    ytd_counts: Counter[str] = Counter()
    ytd_codes: set[str] = set()
    for ledger in ledgers:
        for code, amount in ledger.amount_by_code_yuan.items():
            ytd_amount_yuan[code] += amount
            ytd_codes.add(code)
        ytd_counts.update(ledger.record_count_by_code)

    details: list[dict[str, object]] = []
    for row in report_rows:
        code = str(row["project_code"])
        report_value = row["report_value_wan"]
        report_numeric = 0.0 if report_value is None else float(report_value)
        aug_has_record = current_ledger.record_count_by_code[code] > 0
        ytd_has_record = ytd_counts[code] > 0
        aug_value = near_zero(current_ledger.amount_by_code_yuan.get(code, 0.0) / YUAN_PER_WAN)
        ytd_value = near_zero(ytd_amount_yuan.get(code, 0.0) / YUAN_PER_WAN)
        diff_aug = report_numeric - aug_value
        diff_ytd = report_numeric - ytd_value
        report_blank = bool(row["report_is_blank"])
        details.append(
            {
                "report_row": row["report_row"],
                "project_code": code,
                "project_name": row["project_name"],
                "region": row["region"],
                "line": row["line"],
                "report_is_blank": report_blank,
                "report_value_wan": report_value,
                "in_project_query": code in query_codes,
                "aug_has_eligible_record": aug_has_record,
                "aug_eligible_record_count": current_ledger.record_count_by_code[code],
                "calc_aug_wan": aug_value,
                "diff_aug_wan": diff_aug,
                "match_aug_numeric": report_value is not None and abs(diff_aug) <= TOLERANCE_WAN,
                "match_aug_presence": report_blank == (not aug_has_record),
                "ytd_has_eligible_record": ytd_has_record,
                "ytd_eligible_record_count": ytd_counts[code],
                "calc_ytd_wan": ytd_value,
                "diff_ytd_wan": diff_ytd,
                "match_ytd_numeric": report_value is not None and abs(diff_ytd) <= TOLERANCE_WAN,
                "match_ytd_presence": report_blank == (not ytd_has_record),
            }
        )

    aug_summary = candidate_summary(details, "aug")
    ytd_summary = candidate_summary(details, "ytd")
    if aug_summary["numeric_mismatch_rows"] < ytd_summary["numeric_mismatch_rows"]:
        inferred_basis = "report_month_only"
    elif ytd_summary["numeric_mismatch_rows"] < aug_summary["numeric_mismatch_rows"]:
        inferred_basis = "year_to_date_sum"
    elif aug_summary["max_abs_diff_wan"] < ytd_summary["max_abs_diff_wan"]:
        inferred_basis = "report_month_only"
    elif ytd_summary["max_abs_diff_wan"] < aug_summary["max_abs_diff_wan"]:
        inferred_basis = "year_to_date_sum"
    else:
        inferred_basis = "indeterminate"

    output_dir = ROOT / "output" / "issue-68"
    detail_path = output_dir / f"current_year_concession_{REPORT_MONTH}_project_detail.csv"
    summary_path = output_dir / f"current_year_concession_{REPORT_MONTH}_project_summary.json"
    write_csv(detail_path, details)

    ledger_months = []
    for ledger in ledgers:
        ledger_months.append(
            {
                "month": ledger.month,
                "eligible_rows": ledger.eligible_rows,
                "eligible_project_codes": len(ledger.amount_by_code_yuan),
                "projects_with_multiple_eligible_rows": sum(
                    count > 1 for count in ledger.record_count_by_code.values()
                ),
                "exact_duplicate_rows": ledger.exact_duplicate_rows,
                "blank_code_rows": ledger.blank_code_rows,
                "invalid_amount_rows": ledger.invalid_amount_rows,
                "eligible_amount_total_yuan": sum(ledger.amount_by_code_yuan.values()),
            }
        )

    summary = {
        "report_month": REPORT_MONTH,
        "metric": METRIC,
        "report_unit": u(r"\u4e07\u5143"),
        "ledger_unit": u(r"\u5143"),
        "tolerance_wan": TOLERANCE_WAN,
        "source_files": {
            "report": report_path.name,
            "ledger_month_count": len(ledgers),
            "project_query": query_path.name,
        },
        "report": {
            **report_diagnostics,
            "nonzero_metric_rows": sum(
                row["report_value_wan"] is not None
                and abs(float(row["report_value_wan"])) > TOLERANCE_WAN
                for row in report_rows
            ),
            "codes_missing_from_project_query": sum(
                str(row["project_code"]) not in query_codes for row in report_rows
            ),
        },
        "project_query": query_diagnostics,
        "ledger_months": ledger_months,
        "current_month_candidate": aug_summary,
        "year_to_date_candidate": ytd_summary,
        "ledger_codes_not_in_report": {
            "current_month_count": len(set(current_ledger.amount_by_code_yuan) - report_codes),
            "current_month_amount_yuan": sum(
                amount
                for code, amount in current_ledger.amount_by_code_yuan.items()
                if code not in report_codes
            ),
            "year_to_date_count": len(ytd_codes - report_codes),
            "year_to_date_amount_yuan": sum(
                amount for code, amount in ytd_amount_yuan.items() if code not in report_codes
            ),
        },
        "inferred_report_basis": inferred_basis,
        "outputs": {
            "detail_csv": str(detail_path.relative_to(ROOT)),
            "summary_json": str(summary_path.relative_to(ROOT)),
        },
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
